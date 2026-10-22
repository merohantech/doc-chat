"""Turn an uploaded file into plain text for the model.

Extraction runs locally (no per-request parsing fees) and produces the same
text every time, so a file sent once stays byte-identical in the cached prompt.
The file type is decided from the extension *and* checked against the file's
leading bytes; the client's Content-Type header is not trusted.
"""

import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import PurePath

import docx
from pypdf import PdfReader
from pypdf.errors import PdfReadError

SUPPORTED_TYPES: dict[str, str] = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".csv": "text/csv",
}

# Below this many non-whitespace characters per page, a PDF is almost certainly
# scanned images rather than text.
_MIN_CHARS_PER_PDF_PAGE = 30


class ExtractionError(Exception):
    code = "file_unreadable"

    def __init__(self, message: str) -> None:
        super().__init__(message)
        self.message = message


class UnsupportedFileTypeError(ExtractionError):
    code = "file_type_unsupported"


@dataclass(frozen=True, slots=True)
class ExtractedText:
    text: str
    mime: str


def extract_text(filename: str, data: bytes) -> ExtractedText:
    extension = PurePath(filename).suffix.lower()
    mime = SUPPORTED_TYPES.get(extension)
    if mime is None:
        supported = ", ".join(sorted(SUPPORTED_TYPES))
        raise UnsupportedFileTypeError(f"Unsupported file type. Supported types: {supported}.")

    if extension == ".pdf":
        text = _pdf(data)
    elif extension == ".docx":
        text = _docx(data)
    else:
        text = _plain(data)

    text = _normalize(text)
    if not text:
        raise ExtractionError("The file contains no text.")
    return ExtractedText(text=text, mime=mime)


def _pdf(data: bytes) -> str:
    if not data.startswith(b"%PDF"):
        raise ExtractionError("The file is not a valid PDF.")
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractionError("The PDF is password-protected.")
        pages = [page.extract_text() or "" for page in reader.pages]
    except PdfReadError as exc:
        raise ExtractionError("The PDF could not be read.") from exc

    visible = sum(len("".join(page.split())) for page in pages)
    if visible < _MIN_CHARS_PER_PDF_PAGE * max(len(pages), 1):
        raise ExtractionError(
            "This PDF appears to be scanned images with no text layer. "
            "Upload a text-based PDF instead."
        )
    # Page markers let the model cite pages ("see page 3").
    return "\n\n".join(f"--- Page {i} ---\n{page.strip()}" for i, page in enumerate(pages, 1))


def _docx(data: bytes) -> str:
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise ExtractionError("The file is not a valid Word document.")
    try:
        document = docx.Document(io.BytesIO(data))
    except (KeyError, ValueError, zipfile.BadZipFile) as exc:
        raise ExtractionError("The Word document could not be read.") from exc

    blocks = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(blocks)


def _plain(data: bytes) -> str:
    if b"\x00" in data:
        raise ExtractionError("The file looks binary, not text.")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Common for CSVs exported from Excel on Windows.
        return data.decode("cp1252", errors="replace")


def _normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    return re.sub(r"\n{3,}", "\n\n", text).strip()
