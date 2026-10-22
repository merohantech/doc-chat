import io

import docx
import pytest
from fpdf import FPDF

from app.ingestion.extractors import (
    ExtractionError,
    UnsupportedFileTypeError,
    extract_text,
)


def make_pdf(*pages: str) -> bytes:
    pdf = FPDF()
    pdf.set_font("Helvetica", size=12)
    for text in pages:
        pdf.add_page()
        if text:
            pdf.multi_cell(0, 10, text)
    return bytes(pdf.output())


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_pdf_text_has_page_markers() -> None:
    data = make_pdf(
        "Quarterly revenue grew twelve percent year over year.",
        "Headcount stayed flat at forty two people across all teams.",
    )

    result = extract_text("report.pdf", data)

    assert result.mime == "application/pdf"
    assert result.text.startswith("--- Page 1 ---\nQuarterly revenue grew")
    assert "--- Page 2 ---\nHeadcount stayed flat" in result.text


def test_scanned_pdf_without_text_is_rejected() -> None:
    with pytest.raises(ExtractionError, match="scanned"):
        extract_text("scan.pdf", make_pdf("", ""))


def test_non_pdf_bytes_with_pdf_extension_are_rejected() -> None:
    with pytest.raises(ExtractionError, match="not a valid PDF"):
        extract_text("fake.pdf", b"hello, I am not a PDF")


def test_docx_paragraphs_and_tables() -> None:
    data = make_docx(
        ["Project plan", "Phase one starts in May."], [["Owner", "Task"], ["Ada", "API"]]
    )

    result = extract_text("plan.docx", data)

    assert result.text == "Project plan\nPhase one starts in May.\nOwner | Task\nAda | API"


def test_non_zip_docx_is_rejected() -> None:
    with pytest.raises(ExtractionError, match="not a valid Word document"):
        extract_text("plan.docx", b"plain text pretending to be docx")


@pytest.mark.parametrize("name", ["notes.txt", "README.md", "data.csv", "NOTES.TXT"])
def test_text_formats_are_decoded(name: str) -> None:
    result = extract_text(name, "name,city\r\nAda,Zürich\r\n".encode())

    assert result.text == "name,city\nAda,Zürich"


def test_windows_1252_text_falls_back_cleanly() -> None:
    result = extract_text("legacy.csv", "café;prix\n".encode("cp1252"))

    assert result.text == "café;prix"


def test_excess_blank_lines_are_collapsed() -> None:
    result = extract_text("a.txt", b"one\n\n\n\n\ntwo   \n")

    assert result.text == "one\n\ntwo"


def test_binary_content_is_rejected() -> None:
    with pytest.raises(ExtractionError, match="binary"):
        extract_text("data.txt", b"\x00\x01\x02")


def test_empty_file_is_rejected() -> None:
    with pytest.raises(ExtractionError, match="no text"):
        extract_text("empty.txt", b"   \n  ")


def test_unsupported_extension() -> None:
    with pytest.raises(UnsupportedFileTypeError) as excinfo:
        extract_text("photo.png", b"\x89PNG")
    assert excinfo.value.code == "file_type_unsupported"
